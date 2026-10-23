# agent/informes.py — Informes descargables del panel

"""
Arma el Excel del resumen semanal (agent/memory.py: obtener_resumen_semanal) para el boton
"Descargar" del dashboard. Dos hojas: el resumen con lo que significa cada numero, y el
detalle dia por dia.
"""

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

_AZUL = "3B6FE0"
_BORDE = Side(style="thin", color="DFE1E8")
_DIAS = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]


def _fecha(iso: str) -> str:
    d = date.fromisoformat(iso)
    return f"{d:%d/%m/%Y}"


def texto_resumen(r: dict) -> str:
    """El resumen en una frase, listo para pegar en un Word, un mail o un WhatsApp."""
    texto = (
        f"Resumen semanal de PAMPA, del {_fecha(r['desde'])} al {_fecha(r['hasta'])}: "
        f"escribieron {r['personas_que_escribieron']} personas. La IA atendió a "
        f"{r['atendidas_por_ia']} y se derivaron {r['derivadas_a_persona']} a una persona del "
        f"equipo ({r['porcentaje_derivadas']:.1f}% de derivación). "
        f"{r['atendidas_solo_por_ia']} personas se atendieron solo con la IA, sin derivar. "
        f"Mensajes recibidos: {r['mensajes_recibidos']}; respuestas de la IA: {r['respuestas_de_ia']}."
    )
    if r["derivadas_aproximado"]:
        texto += " (El dato de derivaciones es aproximado para esta semana.)"
    return texto


def resumen_a_xlsx(r: dict) -> bytes:
    libro = Workbook()

    # ── Hoja 1: resumen ──────────────────────────────────────────────────────
    hoja = libro.active
    hoja.title = "Resumen"
    hoja["A1"] = "PAMPA — Resumen semanal"
    hoja["A1"].font = Font(bold=True, size=16, color=_AZUL)
    hoja["A2"] = f"Semana del {_fecha(r['desde'])} al {_fecha(r['hasta'])} (lunes a domingo, hora de Argentina)"
    hoja["A2"].font = Font(color="8A8D99")

    filas = [
        ("Personas que escribieron", r["personas_que_escribieron"], "Personas distintas que mandaron al menos un mensaje en la semana."),
        ("Atendidas por la IA", r["atendidas_por_ia"], "De ellas, las que recibieron al menos una respuesta de la IA."),
        ("Derivadas a una persona", r["derivadas_a_persona"], "Personas para las que la IA avisó al equipo porque necesitaban atención humana."),
        ("% de derivación", r["porcentaje_derivadas"] / 100, "Derivadas sobre el total de personas de la semana."),
        ("Atendidas solo por la IA", r["atendidas_solo_por_ia"], "Atendidas por la IA que no hizo falta derivar."),
        ("Con respuesta del equipo", r["con_respuesta_del_equipo"], "Personas a las que alguien del equipo les contestó a mano desde el panel."),
        ("Contactos nuevos", r["contactos_nuevos"], "Números que escribieron por primera vez en la semana."),
        ("Mensajes recibidos", r["mensajes_recibidos"], "Mensajes de clientes."),
        ("Respuestas de la IA", r["respuestas_de_ia"], "Mensajes enviados por la IA."),
        ("Respuestas del equipo", r["respuestas_del_equipo"], "Mensajes enviados a mano por el equipo desde el panel."),
    ]
    encabezado = ["Indicador", "Valor", "Qué significa"]
    for col, titulo in enumerate(encabezado, start=1):
        c = hoja.cell(row=4, column=col, value=titulo)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=_AZUL)
        c.alignment = Alignment(vertical="center")
    for i, (nombre, valor, explicacion) in enumerate(filas, start=5):
        hoja.cell(row=i, column=1, value=nombre).font = Font(bold=True)
        v = hoja.cell(row=i, column=2, value=valor)
        v.alignment = Alignment(horizontal="right")
        if nombre.startswith("%"):
            v.number_format = "0.0%"
        hoja.cell(row=i, column=3, value=explicacion).font = Font(color="8A8D99")
        for col in (1, 2, 3):
            hoja.cell(row=i, column=col).border = Border(bottom=_BORDE)
    if r["derivadas_aproximado"]:
        nota = hoja.cell(
            row=len(filas) + 6,
            column=1,
            value="Nota: las derivaciones de esta semana son aproximadas (se registran una por una recién desde que se activó el seguimiento).",
        )
        nota.font = Font(italic=True, color="B3760F")
    hoja.column_dimensions["A"].width = 30
    hoja.column_dimensions["B"].width = 12
    hoja.column_dimensions["C"].width = 78

    # ── Hoja 2: por dia ──────────────────────────────────────────────────────
    detalle = libro.create_sheet("Por día")
    columnas = ["Día", "Fecha", "Personas que escribieron", "Mensajes recibidos", "Respuestas de la IA", "Derivadas a una persona"]
    for col, titulo in enumerate(columnas, start=1):
        c = detalle.cell(row=1, column=col, value=titulo)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=_AZUL)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, d in enumerate(r["por_dia"], start=2):
        dia = date.fromisoformat(d["dia"])
        valores = [_DIAS[dia.weekday()], f"{dia:%d/%m/%Y}", d["personas"], d["mensajes_cliente"], d["respuestas_ia"], d["derivadas"]]
        for col, valor in enumerate(valores, start=1):
            c = detalle.cell(row=i, column=col, value=valor)
            c.border = Border(bottom=_BORDE)
            if col >= 3:
                c.alignment = Alignment(horizontal="right")
    detalle.row_dimensions[1].height = 32
    for letra, ancho in zip("ABCDEF", (14, 13, 16, 16, 16, 16)):
        detalle.column_dimensions[letra].width = ancho

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
